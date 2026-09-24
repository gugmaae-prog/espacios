from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "HANDOVER.md"
OUTPUT = ROOT / "ESPACIOS_MAP_FULLSTACK_HANDOVER_2026-09-22.docx"

INK = "000000"
MUTED = "586174"
NAVY = "172033"
PALE = "F3F6FA"
PALE_BLUE = "EDF2FA"
BORDER = "D9D9D9"
CODE_BG = "F5F6F8"


def set_run_font(run, name: str = "Arial", size: float | None = None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "5")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), BORDER)


def remove_paragraph_borders(paragraph_or_style):
    element = paragraph_or_style._element
    if hasattr(element, "get_or_add_pPr"):
        p_pr = element.get_or_add_pPr()
    else:
        p_pr = paragraph_or_style._p.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def add_field(run, instruction: str):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))


def normalize_text(text: str) -> str:
    return (
        text.replace("—", "-")
        .replace("–", "-")
        .replace("→", "->")
        .replace("“", '"')
        .replace("”", '"')
        .replace("’", "'")
        .replace(" ", " ")
    )


INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|<https?://[^>]+>)")


def add_inline(paragraph, text: str, *, base_size: float = 10.5):
    text = normalize_text(text.strip())
    for token in INLINE_RE.split(text):
        if not token:
            continue
        run = paragraph.add_run()
        if token.startswith("**") and token.endswith("**"):
            run.text = token[2:-2]
            run.bold = True
            set_run_font(run, size=base_size)
        elif token.startswith("`") and token.endswith("`"):
            run.text = token[1:-1]
            set_run_font(run, "Courier New", max(9.0, base_size - 0.5))
        elif token.startswith("<http") and token.endswith(">"):
            run.text = token[1:-1]
            run.font.color.rgb = RGBColor.from_string("2449A8")
            run.font.underline = True
            set_run_font(run, size=base_size)
        else:
            run.text = token
            set_run_font(run, size=base_size)


def clean_heading(text: str) -> str:
    text = normalize_text(text)
    text = re.sub(r"^(\d+)\.\s+", r"\1 ", text)
    text = text.replace("/", " and ")
    text = text.replace(":", "")
    text = text.replace("?", "")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def style_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.widow_control = True

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(INK)
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.keep_with_next = True
    remove_paragraph_borders(title)

    for name, size, before, after in (
        ("Heading 1", 17, 16, 7),
        ("Heading 2", 13.5, 12, 5),
        ("Heading 3", 11.5, 9, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(10.3)
        style.paragraph_format.left_indent = Inches(0.24)
        style.paragraph_format.first_line_indent = Inches(-0.16)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.08

    code = doc.styles.add_style("Espacios Code", 1)
    code.font.name = "Courier New"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Courier New")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")
    code.font.size = Pt(8.8)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.12)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(7)
    code.paragraph_format.line_spacing = 1.0


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run("Espacios UAE Intelligence Map  |  Full Stack Handover  |  ")
    r.font.color.rgb = RGBColor.from_string(MUTED)
    set_run_font(r, size=8)
    page = p.add_run()
    page.font.color.rgb = RGBColor.from_string(MUTED)
    set_run_font(page, size=8)
    add_field(page, "PAGE")
    r2 = p.add_run(" of ")
    r2.font.color.rgb = RGBColor.from_string(MUTED)
    set_run_font(r2, size=8)
    total = p.add_run()
    total.font.color.rgb = RGBColor.from_string(MUTED)
    set_run_font(total, size=8)
    add_field(total, "NUMPAGES")


def add_metadata_table(doc: Document):
    values = [
        ("Verified", "22 September 2026  Asia Dubai"),
        ("Production URL", "https://espacios.me/map"),
        ("Worker", "psr-portfolio-map-v2"),
        ("Release", "20260921-seamless-v12"),
        ("Purpose", "Production architecture release operations and rollback handover"),
    ]
    table = doc.add_table(rows=len(values), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.45)
    table.columns[1].width = Inches(5.2)
    set_table_borders(table)
    for index, (label, value) in enumerate(values):
        for cell in table.rows[index].cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 130, 140, 130, 140)
            if index % 2:
                shade_cell(cell, PALE)
        p1 = table.cell(index, 0).paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(label)
        r1.bold = True
        set_run_font(r1, size=9.5)
        p2 = table.cell(index, 1).paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        add_inline(p2, value, base_size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_markdown_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    if doc.paragraphs:
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    repeat_table_header(table.rows[0])

    if cols == 2:
        widths = (2.05, 4.6)
    elif cols == 3:
        widths = (1.7, 1.45, 3.5)
    elif cols == 4:
        widths = (1.25, 1.8, 1.8, 1.8)
    else:
        widths = tuple(6.65 / cols for _ in range(cols))

    for r_idx, source_row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.width = Inches(widths[c_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 95, 105, 95, 105)
            text = source_row[c_idx] if c_idx < len(source_row) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.03
            if r_idx == 0:
                shade_cell(cell, NAVY)
                r = p.add_run(normalize_text(text))
                r.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                set_run_font(r, size=9.2)
            else:
                if r_idx % 2 == 0:
                    shade_cell(cell, PALE_BLUE)
                add_inline(p, text, base_size=8.9)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)


def parse_table(lines: list[str], start: int):
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def build_document():
    source_lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style_document(doc)
    add_footer(doc.sections[0])

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    remove_paragraph_borders(title)
    add_inline(title, "Espacios UAE Intelligence Map Full Stack Handover", base_size=27)
    for run in title.runs:
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(INK)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    r = subtitle.add_run("Current production architecture release operations and rollback")
    r.font.color.rgb = RGBColor.from_string(MUTED)
    set_run_font(r, size=13)

    add_metadata_table(doc)

    purpose = doc.add_paragraph(style="Heading 1")
    purpose.add_run("Purpose")
    intro = doc.add_paragraph()
    add_inline(
        intro,
        "This document is the canonical operational handover for the Espacios UAE Intelligence Map. It records what is live, how the system is assembled, how to work locally, how to release safely, how to verify production, and how to roll back. It supersedes the older v10 release metadata retained in the project.",
    )
    doc.add_page_break()

    start = next(i for i, line in enumerate(source_lines) if line.startswith("## 1."))
    lines = source_lines[start:]
    index = 0
    in_code = False
    code_lines: list[str] = []

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()

        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph(style="Espacios Code")
                p.paragraph_format.keep_together = True
                p_pr = p._p.get_or_add_pPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), CODE_BG)
                p_pr.append(shd)
                r = p.add_run(normalize_text("\n".join(code_lines)))
                set_run_font(r, "Courier New", 8.8)
                in_code = False
            index += 1
            continue

        if in_code:
            code_lines.append(raw)
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            add_markdown_table(doc, rows)
            continue

        heading = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading:
            level = min(3, len(heading.group(1)) - 1)
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline(p, clean_heading(heading.group(2)), base_size={1: 17, 2: 13.5, 3: 11.5}[level])
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(INK)
            index += 1
            continue

        bullet = re.match(r"^-\s+(.+)$", stripped)
        if bullet:
            text = bullet.group(1)
            if text.startswith("[ ] "):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.28)
                p.paragraph_format.first_line_indent = Inches(-0.18)
                p.paragraph_format.space_after = Pt(3)
                marker = p.add_run("☐  ")
                set_run_font(marker, "Segoe UI Symbol", 10.5)
                add_inline(p, text[4:], base_size=10.3)
            else:
                p = doc.add_paragraph(style="List Bullet")
                add_inline(p, text, base_size=10.3)
            index += 1
            continue

        numbered = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if numbered:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            p.paragraph_format.space_after = Pt(3)
            marker = p.add_run(numbered.group(1) + ". ")
            set_run_font(marker, size=10.3)
            add_inline(p, numbered.group(2), base_size=10.3)
            index += 1
            continue

        parts = [stripped]
        index += 1
        while index < len(lines):
            nxt = lines[index].strip()
            if (
                not nxt
                or nxt.startswith("#")
                or nxt.startswith("|")
                or nxt.startswith("```")
                or re.match(r"^-\s+", nxt)
                or re.match(r"^\d+\.\s+", nxt)
            ):
                break
            parts.append(nxt)
            index += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(parts))

    props = doc.core_properties
    props.title = "Espacios UAE Intelligence Map Full Stack Handover"
    props.subject = "Production architecture release operations verification and rollback"
    props.author = "Espacios"
    props.keywords = "Espacios UAE map Cloudflare Worker R2 D1 MapLibre handover"
    props.comments = "Verified production handover dated 22 September 2026"

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
