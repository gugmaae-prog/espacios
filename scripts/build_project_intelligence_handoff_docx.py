"""Build the 24 Sep 2026 project-level intelligence handoff DOCX from markdown."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))

from build_handover_docx import (  # noqa: E402
    CODE_BG,
    INK,
    MUTED,
    add_footer,
    add_inline,
    add_markdown_table,
    clean_heading,
    normalize_text,
    parse_table,
    set_run_font,
    style_document,
)


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "ESPACIOS_FULLSTACK_HANDOFF_PROJECT_LEVEL_INTELLIGENCE_2026-09-24.md"
OUTPUT = ROOT / "ESPACIOS_FULLSTACK_HANDOFF_PROJECT_LEVEL_INTELLIGENCE_2026-09-24.docx"


def add_metadata_table(doc: Document):
    from build_handover_docx import (
        PALE,
        WD_CELL_VERTICAL_ALIGNMENT,
        WD_TABLE_ALIGNMENT,
        set_cell_margins,
        set_table_borders,
        shade_cell,
    )

    values = [
        ("Checked", "24 September 2026"),
        ("Production URL", "https://espacios.me/map"),
        ("Worker", "psr-portfolio-map-v2"),
        ("Live version", "145 at 100% traffic"),
        ("Version UUID", "bb5456b0-d558-4dba-9856-a80e19f4b5a2"),
        ("Deployment UUID", "01fa63ce-ab72-4592-99aa-b211ec857dc9"),
        ("Workbook", "1,382 projects / 528 scenario-ready / 0 approved forecasts"),
        ("Purpose", "Project-level pricing, sourcing, APIs, and release boundaries"),
    ]
    table = doc.add_table(rows=len(values), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.55)
    table.columns[1].width = Inches(5.1)
    set_table_borders(table)
    for index, (label, value) in enumerate(values):
        for cell in table.rows[index].cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 110, 130, 110, 130)
            if index % 2:
                shade_cell(cell, PALE)
        p1 = table.cell(index, 0).paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(label)
        r1.bold = True
        set_run_font(r1, size=9.5)
        p2 = table.cell(index, 1).paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        add_inline(p2, value, base_size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build_document():
    source_lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style_document(doc)
    add_footer(doc.sections[0])

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_inline(title, "Espacios Project-Level Intelligence Handoff", base_size=26)
    for run in title.runs:
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(INK)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    r = subtitle.add_run(
        "Historical pricing, scenario modelling, source governance, and deployment boundaries"
    )
    r.font.color.rgb = RGBColor.from_string(MUTED)
    set_run_font(r, size=12)

    add_metadata_table(doc)
    doc.add_page_break()

    start = next(i for i, line in enumerate(source_lines) if line.startswith("## Contents"))
    lines = source_lines[start:]
    index = 0
    in_code = False
    code_lines: list[str] = []

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()

        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph(style="Espacios Code")
                p.paragraph_format.keep_together = True
                p_pr = p._p.get_or_add_pPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), CODE_BG)
                p_pr.append(shd)
                run = p.add_run(normalize_text("\n".join(code_lines)))
                set_run_font(run, "Courier New", 8.8)
                in_code = False
            index += 1
            continue

        if in_code:
            code_lines.append(raw)
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            add_markdown_table(doc, rows)
            continue

        heading = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading:
            level = min(3, len(heading.group(1)) - 1)
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline(p, clean_heading(heading.group(2)), base_size={1: 17, 2: 13.5, 3: 11.5}[level])
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(INK)
            index += 1
            continue

        bullet = re.match(r"^-\s+(.+)$", stripped)
        if bullet:
            text = bullet.group(1)
            if text.startswith("[ ] "):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.28)
                p.paragraph_format.first_line_indent = Inches(-0.18)
                p.paragraph_format.space_after = Pt(3)
                marker = p.add_run("[ ]  ")
                set_run_font(marker, size=10.5)
                add_inline(p, text[4:], base_size=10.3)
            else:
                p = doc.add_paragraph(style="List Bullet")
                add_inline(p, text, base_size=10.3)
            index += 1
            continue

        numbered = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if numbered:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            p.paragraph_format.space_after = Pt(3)
            marker = p.add_run(numbered.group(1) + ". ")
            set_run_font(marker, size=10.3)
            add_inline(p, numbered.group(2), base_size=10.3)
            index += 1
            continue

        parts = [stripped]
        index += 1
        while index < len(lines):
            nxt = lines[index].strip()
            if (
                not nxt
                or nxt.startswith("#")
                or nxt.startswith("|")
                or nxt.startswith("```")
                or re.match(r"^-\s+", nxt)
                or re.match(r"^\d+\.\s+", nxt)
            ):
                break
            parts.append(nxt)
            index += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(parts))

    props = doc.core_properties
    props.title = "Espacios Project-Level Intelligence Full-Stack Handoff"
    props.subject = "Project-level historical pricing, sourcing, scenarios, and deployment boundaries"
    props.author = "Espacios"
    props.category = "Engineering handoff"
    props.keywords = "Espacios UAE project intelligence DLD sourcing scenarios Cloudflare"
    props.comments = "Checked against production Worker version 145 on 24 September 2026. Workbook scenarios are not deployed."

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
